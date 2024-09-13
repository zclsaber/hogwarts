import os
import pytesseract
from PIL import Image
from openpyxl import load_workbook

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

img_path1 = r"C:\Users\zcl\Downloads\27_2023-10-31\20231031053907_FS_8cLkq6sw\2\Screenshot_2023-10-27-22-13-12-36_6a0285d7878ab47a2b69dc50bfc0de_207871905.jpg"
img_path2 = r"C:\Users\zcl\Downloads\27_2023-10-31\20231031053907_FS_8cLkq6sw\1\mmexport1698576101273_207871905.jpg"
img_path3 = r"C:\Users\zcl\Downloads\27_2023-10-31\20231031053907_FS_8cLkq6sw\1\mmexport1698576111296_207871904.jpg"

class FileHandler:
    '''
    处理从麦客导出的整体目录，将整理好后的信息标准化输出
    '''
    def __int__(self):
        pass

    def directory_process(self):
        '''
        遍历目录，输出一个目录结构的字典，类似于{"1": [img1, img2, img3], "2": [img1, img2, img3], ...}
        :return:
        '''
        pass

    def file_process(self, file):
        '''
        对导出的麦客 excel 文件进行操作，在对应的行末尾插入所需要的值
        先判断列是否存在，如果列不存在，则添加列；如果列存在，则直接在后面添加数据
        需要的列： 1. 是否提交； 2. 是否是对应日期的票； 3. 实际金额是多少； 4. 最终报销的金额是多少； 5. 是否 10 星好评
        :return:
        '''
        workbook = load_workbook(file)
        worksheet = workbook['Sheet1']

        # 先遍历列，之后确认列是否存在
        header_row = worksheet[1]
        header = [cell.value for cell in header_row]

        header_for_add = ["是否提交", "对应日期", "实际金额", "报销金额", "10星好评"]

        added_column = 0
        for column in header_for_add:
            if column not in header:
                added_column += 1
                worksheet.cell(row=1, colunm=len(header)+added_column, value=column)

        workbook.save(file)
        workbook.close()

class ImgProcessor:
    '''
    处理单个文件夹下的内容，也即是处理一个报名用户的信息
    '''
    def __int__(self):
        pass

    def excel_editor(self):
        pass



    def check_review(self):
        pass

    def check_payment(self):
        pass

    def check_seat(self):
        pass

    def check_actual_payment(self):
        pass


img = Image.open(img_path1)

txt = pytesseract.image_to_string(img, lang='chi_sim')
print(txt)